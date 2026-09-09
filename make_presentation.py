import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation with 16:9 Widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6] # Blank slide layout

# Professional Color Palette
BG_COLOR = RGBColor(11, 22, 44)         # Deep Navy #0b162c
CARD_BG = RGBColor(18, 35, 68)          # Card Blue #122344
CARD_BORDER = RGBColor(38, 70, 120)     # Border Blue #264678
ACCENT_BLUE = RGBColor(0, 132, 255)     # Brand Blue #0084ff
ACCENT_CYAN = RGBColor(56, 189, 248)    # Cyan #38bdf8
ACCENT_AMBER = RGBColor(245, 158, 11)   # Amber #f59e0b
ACCENT_GREEN = RGBColor(16, 185, 129)   # Emerald #10b981
TEXT_WHITE = RGBColor(248, 250, 252)    # Slate-50 #f8fafc
TEXT_MUTED = RGBColor(148, 163, 184)    # Slate-400 #94a3b8
TEXT_LIGHT = RGBColor(203, 213, 225)    # Slate-300 #cbd5e1

def set_slide_background(slide):
    # Full slide background shape
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

def add_header(slide, category_text, title_text):
    # Category Pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(4.5), Inches(0.38))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(15, 30, 60)
    pill.line.color.rgb = ACCENT_CYAN
    pill.line.width = Pt(1)
    
    tf = pill.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = category_text.upper()
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.CENTER
    
    # Title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.88), Inches(11.733), Inches(0.65))
    tf2 = tb.text_frame
    tf2.word_wrap = True
    tf2.vertical_anchor = MSO_ANCHOR.TOP
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.font.name = "Arial"

# ==========================================
# SLIDE 1: Title & Technology Overview
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

# Decorative Glow Card Top
top_banner = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.42))
top_banner.fill.solid()
top_banner.fill.fore_color.rgb = RGBColor(14, 30, 62)
top_banner.line.color.rgb = ACCENT_BLUE
top_banner.line.width = Pt(1)
tf = top_banner.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.text = "SMART INDIA HACKATHON (SIH) FORMAT  •  WEB SCRAPING & DATA ANALYSIS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p.font.name = "Arial"
p.alignment = PP_ALIGN.CENTER

# Main Title Box
title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.733), Inches(1.8))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Naukri.com IT Job Vacancy &\nSalary Trend Analyzer"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = TEXT_WHITE
p.font.name = "Arial"
p.alignment = PP_ALIGN.LEFT

p2 = tf.add_paragraph()
p2.text = "Automated Web Scraping, Compensation Analytics & Career Intelligence Platform"
p2.font.size = Pt(14)
p2.font.color.rgb = ACCENT_AMBER
p2.font.bold = True
p2.font.name = "Arial"
p2.space_before = Pt(8)

# 4 Key Tech Stack Pillars
tech_pillars = [
    ("WEB SCRAPING & INGESTION", "Requests + BeautifulSoup4 + Selenium (Headless Chrome)\nDynamic DOM extraction & live progress tracking", ACCENT_CYAN),
    ("DATA NORMALIZATION & STATS", "Pandas + NumPy + Regex (re)\nUnstructured text to numeric LPA & experience floats", ACCENT_GREEN),
    ("DATA VISUALIZATION", "Plotly.js + Matplotlib / Plotly Python SDK\n8 Interactive charts: Bar, Donut, Histograms, Scatter", ACCENT_AMBER),
    ("DATABASE & FULL-STACK UX", "SQLite3 (B-Tree Indexed) + Flask + Tailwind CSS\nREST APIs, CRUD administration, CSV/Excel/PDF exports", ACCENT_BLUE)
]

card_w = Inches(2.78)
gap = Inches(0.2)
start_x = Inches(0.8)
start_y = Inches(3.3)
card_h = Inches(3.5)

for i, (head, desc, color) in enumerate(tech_pillars):
    x = start_x + i * (card_w + gap)
    card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, start_y, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = color
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.25)
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    
    p = tf.paragraphs[0]
    p.text = head
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = color
    p.font.name = "Arial"
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_LIGHT
    p2.font.name = "Arial"
    p2.space_before = Pt(10)

# ==========================================
# SLIDE 2: Problem Statement & Proposed Solution
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_header(slide2, "Domain: Web Scraping & Data Analysis", "Problem Statement & Proposed Solution")

col_w = Inches(5.7)
col_h = Inches(5.4)
col_y = Inches(1.6)

# Left Box: Problem Statement
prob_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), col_y, col_w, col_h)
prob_card.fill.solid()
prob_card.fill.fore_color.rgb = CARD_BG
prob_card.line.color.rgb = RGBColor(239, 68, 68) # Red-500
prob_card.line.width = Pt(1.5)

tf = prob_card.text_frame
tf.word_wrap = True
tf.margin_top = Inches(0.3)
tf.margin_left = Inches(0.3)
tf.margin_right = Inches(0.3)

p = tf.paragraphs[0]
p.text = "⚠️ THE INDUSTRY PROBLEM"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = RGBColor(248, 113, 113)

problems = [
    ("Salary Opacity & Inconsistency:", "Compensation strings on Naukri are unstructured (e.g. '12-25 Lacs PA', '15,00,000 - 25,00,000 PA', 'Not Disclosed'), preventing standardized comparison."),
    ("Geographic Disparities:", "Major tech hubs (Bengaluru, Hyderabad, Pune, Delhi NCR, Mumbai) feature drastically different hiring volumes and compensation brackets."),
    ("Skill-to-Wage Disconnect:", "Job seekers lack aggregated visibility into which programming languages and cloud stacks command higher salary premiums."),
    ("Manual Inefficiency:", "Manually browsing hundreds of pages on job portals is tedious, unscalable, and prone to subjective bias.")
]

for title, desc in problems:
    p_item = tf.add_paragraph()
    p_item.text = f"• {title} {desc}"
    p_item.font.size = Pt(10.5)
    p_item.font.color.rgb = TEXT_LIGHT
    p_item.space_before = Pt(12)

# Right Box: Proposed Solution
sol_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), col_y, col_w, col_h)
sol_card.fill.solid()
sol_card.fill.fore_color.rgb = CARD_BG
sol_card.line.color.rgb = ACCENT_GREEN
sol_card.line.width = Pt(1.5)

tf2 = sol_card.text_frame
tf2.word_wrap = True
tf2.margin_top = Inches(0.3)
tf2.margin_left = Inches(0.3)
tf2.margin_right = Inches(0.3)

p = tf2.paragraphs[0]
p.text = "💡 OUR PROPOSED SOLUTION"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

solutions = [
    ("Automated Ingestion Engine:", "Headless Selenium + BeautifulSoup4 extracts real-time vacancies, companies, salaries, locations, and tech tags on demand."),
    ("Algorithmic Normalization:", "Custom regex parser standardizes all textual salary strings into numeric LPA (Lakhs Per Annum) and years of experience."),
    ("Executive Career Intelligence:", "LinkedIn-Analytics styled dashboard featuring 8+ interactive Plotly.js charts and live KPI metric summary cards."),
    ("Full Governance & Export:", "SQLite database with CRUD administration and 1-click multi-format report exports (CSV, Excel .xlsx, and PDF).")
]

for title, desc in solutions:
    p_item = tf2.add_paragraph()
    p_item.text = f"✔ {title} {desc}"
    p_item.font.size = Pt(10.5)
    p_item.font.color.rgb = TEXT_LIGHT
    p_item.space_before = Pt(12)

# ==========================================
# SLIDE 3: Technical Architecture & 4-Stage Pipeline
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_header(slide3, "System Design & Methodology", "Technical Architecture & 4-Stage Data Pipeline")

stages = [
    ("STAGE 1: EXTRACTION", "Requests + BS4 + Selenium", "• Headless Chrome engine\n• Dynamic JavaScript scrolling\n• Thread-safe progress tracking\n• Fallback dataset generation", ACCENT_CYAN),
    ("STAGE 2: NORMALIZATION", "Regex + Pandas + NumPy", "• Regex LPA parser algorithm\n• Experience range parsing\n• Location hub standardization\n• Missing value imputation", ACCENT_GREEN),
    ("STAGE 3: PERSISTENCE", "SQLite3 Relational DB", "• 14 Structured table attributes\n• Composite B-Tree indexing\n• Fast indexed search queries\n• 1-Click sample data seeder", ACCENT_AMBER),
    ("STAGE 4: VISUALIZATION", "Plotly.js + Flask + UI", "• 8 Interactive Plotly charts\n• Asynchronous JSON delivery\n• Dual Dark/Light theme mode\n• CSV / Excel / PDF exports", ACCENT_BLUE)
]

stage_w = Inches(2.78)
gap_s = Inches(0.2)
start_x_s = Inches(0.8)
start_y_s = Inches(1.6)
stage_h = Inches(3.2)

for i, (title, tech, bullets, color) in enumerate(stages):
    x = start_x_s + i * (stage_w + gap_s)
    scard = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, start_y_s, stage_w, stage_h)
    scard.fill.solid()
    scard.fill.fore_color.rgb = CARD_BG
    scard.line.color.rgb = color
    scard.line.width = Pt(1.5)
    
    tf = scard.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.2)
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = color
    
    p_tech = tf.add_paragraph()
    p_tech.text = tech
    p_tech.font.size = Pt(9.5)
    p_tech.font.bold = True
    p_tech.font.color.rgb = TEXT_WHITE
    p_tech.space_before = Pt(4)
    
    p_b = tf.add_paragraph()
    p_b.text = bullets
    p_b.font.size = Pt(9.5)
    p_b.font.color.rgb = TEXT_LIGHT
    p_b.space_before = Pt(8)

# Bottom Detail Banner on Slide 3
bot_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(11.733), Inches(2.0))
bot_card.fill.solid()
bot_card.fill.fore_color.rgb = RGBColor(14, 28, 54)
bot_card.line.color.rgb = CARD_BORDER
bot_card.line.width = Pt(1)

tf = bot_card.text_frame
tf.word_wrap = True
tf.margin_top = Inches(0.2)
tf.margin_left = Inches(0.3)
tf.margin_right = Inches(0.3)

p = tf.paragraphs[0]
p.text = "CORE ALGORITHMIC HIGHLIGHT: REGEX SALARY & LOCATION NORMALIZATION"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER

p2 = tf.add_paragraph()
p2.text = "• Salary Parsing: '18-30 Lacs PA' -> min_salary: 18.0 LPA | max_salary: 30.0 LPA | avg_salary: 24.0 LPA (Handles Lacs, Lakhs, Crores, & Numeric INR)\n• Location Normalization: 'Bengaluru/Bangalore' -> 'Bengaluru' | 'Gurgaon/Gurugram/Noida' -> 'Delhi NCR' | Standardizes 8 major tech hubs\n• Performance Optimization: B-Tree indexes on avg_salary, min_exp, location, and company deliver sub-millisecond query performance."
p2.font.size = Pt(10)
p2.font.color.rgb = TEXT_LIGHT
p2.space_before = Pt(6)

# ==========================================
# SLIDE 4: Key Features & System Capabilities
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_header(slide4, "Implementation & Capabilities", "Key Features & System Capabilities")

feat_w = Inches(5.7)
feat_h = Inches(2.55)

features_grid = [
    ("📊 INTERACTIVE EXECUTIVE DASHBOARD", 
     "• Live KPI Metric Cards: Total Jobs, Active Companies, Avg LPA, Top Skill, Remote Share.\n• 8 Plotly.js Visualizations: City Share, Company Ranking, Salary Histogram, Experience Donut, Top Skills, Work Mode Split, Salary vs. Exp Scatter, Role Benchmarks.", 
     Inches(0.8), Inches(1.6), ACCENT_CYAN),
    
    ("🔍 LIVE SCRAPING CONSOLE & LOGS", 
     "• Customizable Search: Target role, location, experience bracket, and page depth.\n• Real-Time Progress Modal: Streaming percentage progress, page status, and scrolling execution logs via thread-safe ScrapingProgress tracker.\n• Built-in fallback dataset engine.", 
     Inches(6.833), Inches(1.6), ACCENT_GREEN),
     
    ("📂 JOB DIRECTORY & MULTI-EXPORTERS", 
     "• Filterable & Paginated Directory: Real-time search by role, company, city, and skills.\n• Multi-Format Data Export: Instant 1-click download in CSV, formatted Excel (.xlsx) workbook, or print-ready executive PDF report with direct apply links.", 
     Inches(0.8), Inches(4.35), ACCENT_AMBER),
     
    ("⚙️ DATABASE CRUD & MODERN UI/UX", 
     "• Database Administration: Add job modal, inline edit modal, single delete, wipe all, and seed dataset.\n• Modern Frontend: Tailwind CSS Glassmorphic design, dual Light/Dark theme with localStorage persistence, mobile slide-out drawer, and 6/6 test pass rate.", 
     Inches(6.833), Inches(4.35), ACCENT_BLUE)
]

for title, content, x, y, color in features_grid:
    fcard = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, feat_w, feat_h)
    fcard.fill.solid()
    fcard.fill.fore_color.rgb = CARD_BG
    fcard.line.color.rgb = color
    fcard.line.width = Pt(1.5)
    
    tf = fcard.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.18)
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = color
    
    p_body = tf.add_paragraph()
    p_body.text = content
    p_body.font.size = Pt(9.5)
    p_body.font.color.rgb = TEXT_LIGHT
    p_body.space_before = Pt(6)

# ==========================================
# SLIDE 5: Innovation, Impact & Future Roadmap
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_header(slide5, "Value Proposition & Future Scope", "Innovation, Impact & Future Roadmap")

col3_w = Inches(3.75)
col3_gap = Inches(0.24)
col3_y = Inches(1.6)
col3_h = Inches(5.4)

pillars = [
    ("🚀 CORE INNOVATIONS", [
        ("Automated LPA Normalization:", "Algorithmic conversion of chaotic, non-standard text strings into standardized quantitative benchmarks."),
        ("Asynchronous Plotly Pipeline:", "Zero-latency JSON chart serialization allowing instant re-rendering across dark and light themes."),
        ("Resilient Fallback Engine:", "Continuous analytics availability even when external portal anti-scraping filters change.")
    ], ACCENT_CYAN),
    
    ("🎯 IMPACT & BENEFICIARIES", [
        ("IT Job Seekers & Students:", "Accurately benchmark market compensation, discover top-paying tech stacks, and target hiring hubs."),
        ("Tech Recruiters & HR Teams:", "Analyze competitor compensation structures and real-time demand for specialized roles."),
        ("Academic Institutions:", "Align university curriculums and training with live market industry requirements.")
    ], ACCENT_GREEN),
    
    ("🔮 FUTURE ROADMAP", [
        ("Phase 1: AI Resume Matcher:", "Integrate LLM embeddings to automatically score and match user resumes against live scraped vacancies."),
        ("Phase 2: Predictive Salary ML:", "Train Scikit-Learn regression models to forecast expected compensation based on candidate tech skill portfolios."),
        ("Phase 3: Multi-Portal Scaling:", "Expand connectors to LinkedIn, Indeed, and Glassdoor for aggregated job intelligence.")
    ], ACCENT_AMBER)
]

for i, (title, items, color) in enumerate(pillars):
    x = Inches(0.8) + i * (col3_w + col3_gap)
    pcard = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, col3_y, col3_w, col3_h)
    pcard.fill.solid()
    pcard.fill.fore_color.rgb = CARD_BG
    pcard.line.color.rgb = color
    pcard.line.width = Pt(1.5)
    
    tf = pcard.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.25)
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.22)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = color
    
    for heading, detail in items:
        p_item = tf.add_paragraph()
        p_item.text = f"• {heading} {detail}"
        p_item.font.size = Pt(9.5)
        p_item.font.color.rgb = TEXT_LIGHT
        p_item.space_before = Pt(14)

# Save presentation
out_path = "Naukri_Job_Analyzer_SIH_5_Slides.pptx"
prs.save(out_path)
print(f"Presentation successfully saved to: {out_path}")
